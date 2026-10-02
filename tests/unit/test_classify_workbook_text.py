"""Names that start like a formula (`=1.jpg`) stay text in the classify workbook.

openpyxl stores text starting with `=` as a formula; Excel then reports a damaged file
and the reader finds an empty cell. Written, edited, read back by `sort` and carried
over by the next `classify`, such a name keeps its exact text.
"""

# classify works on naive local dates, as EXIF writes them.
# ruff: noqa: DTZ001

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import TYPE_CHECKING, Final

from openpyxl import load_workbook

from media_hygiene.classify.carry import carry_over
from media_hygiene.classify.carry_types import CarrySource
from media_hygiene.classify.engine import Scope, classify
from media_hygiene.classify.plan_build import build_plan
from media_hygiene.classify.workbook.prefill import Prefill
from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.salvage import salvage
from media_hygiene.classify.workbook.sheets import EventColumn, FileColumn
from media_hygiene.classify.workbook.writer import write_workbook
from media_hygiene.paths.host_paths import HostPathMapper
from tests.unit.test_classify_workbook import SETTINGS, shot

if TYPE_CHECKING:
    from media_hygiene.classify.plan_file import ClassifyPlan

_FILES, _EVENTS = 3, 2  # sheet positions
_NAME = 4  # the Files sheet: the file name
_LOOSE = 8  # loose shots of one day: an event
_FORMULA, _STRING = "f", "s"
_EVENT_NAME: Final = "=Fête"
_FOLDER: Final = "=Trip/+2019"


def _plan() -> ClassifyPlan:
    """Folders and files whose names start with `=`, `+`, `-` and `@`."""
    files = [
        shot("=Party/=1.jpg", datetime(2019, 8, 1, 10)),
        shot("-2019 trip/+1.jpg", datetime(2019, 8, 1, 11)),
        shot("@home/@a.jpg", datetime(2019, 9, 1, 11)),
    ]
    start = datetime(2016, 7, 14, 10)
    files += [
        shot(f"DCIM/-{index}.jpg", start + timedelta(hours=index))
        for index in range(_LOOSE)
    ]
    classification = classify(files, SETTINGS, Scope())
    return build_plan(classification, SETTINGS, HostPathMapper(Path("/data")))


def _formulas(target: Path) -> list[str]:
    """Every cell of the workbook that holds a formula."""
    book = load_workbook(target)
    return [
        f"{sheet.title}!{cell.coordinate}"
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == _FORMULA
    ]


def _typed(target: Path) -> None:
    """Type text starting with `=` in yellow cells, as Excel keeps it: text."""
    book = load_workbook(target)
    for sheet, column, text in (
        (_EVENTS, EventColumn.NAME, _EVENT_NAME),
        (_FILES, FileColumn.FINAL, _FOLDER),
    ):
        cell = book.worksheets[sheet].cell(2, column)
        cell.value = text
        cell.data_type = _STRING  # a Text-format cell: Excel computes nothing
    book.save(target)


def test_names_are_written_as_text(tmp_path: Path) -> None:
    """No formula anywhere; `sort` reads the workbook back without a difference."""
    plan, target = _plan(), tmp_path / "classify.xlsx"
    write_workbook(plan, target)
    assert not _formulas(target)
    files = load_workbook(target).worksheets[_FILES]
    names = [row[_NAME - 1] for row in files.iter_rows(values_only=True)]
    assert {"=1.jpg", "+1.jpg", "@a.jpg"} <= set(names)
    assert read_edits(target, plan).files == {}


def test_typed_text_survives_sort_and_the_carry_over(tmp_path: Path) -> None:
    """Read by `sort`, salvaged, written again by the next `classify`: same text."""
    plan, target = _plan(), tmp_path / "classify.xlsx"
    write_workbook(plan, target)
    _typed(target)
    edits = read_edits(target, plan)
    assert [edit.name for edit in edits.events.values()] == [_EVENT_NAME]
    assert list(edits.files.values()) == [_FOLDER]
    found = salvage(target)
    carried = carry_over(CarrySource("C:\\classify.xlsx", "", found, plan), plan)
    again = tmp_path / "again.xlsx"
    write_workbook(plan, again, Prefill(carried.edits, carried.notes))
    assert not _formulas(again)
    salvaged = salvage(again).edits
    assert [edit.name for edit in salvaged.events.values()] == [_EVENT_NAME]
    assert list(salvaged.files.values()) == [_FOLDER]
