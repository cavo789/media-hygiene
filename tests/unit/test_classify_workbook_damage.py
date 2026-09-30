"""The classify workbook refuses a file that is not one, or whose structure broke."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.sheets import META_SHEET, MetaKey
from media_hygiene.classify.workbook.writer import write_workbook
from media_hygiene.errors import WorkbookError
from tests.unit.test_classify_workbook import make_plan

if TYPE_CHECKING:
    from pathlib import Path

    from media_hygiene.classify.plan_file import ClassifyPlan


@pytest.fixture(name="written")
def written_workbook(tmp_path: Path) -> tuple[ClassifyPlan, Path]:
    """A plan and its workbook."""
    plan = make_plan()
    target = tmp_path / "classify.xlsx"
    write_workbook(plan, target)
    return plan, target


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path: Path) -> None:
    """Missing, or not a zip: a clear error, not a traceback."""
    plan = make_plan()
    with pytest.raises(WorkbookError, match="cannot be opened"):
        read_edits(tmp_path / "missing.xlsx", plan)
    text = tmp_path / "text.xlsx"
    text.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(WorkbookError, match="cannot be opened"):
        read_edits(text, plan)


def test_a_renamed_sheet_is_refused(written: tuple[ClassifyPlan, Path]) -> None:
    """Sheets are found by the names `_meta` recorded."""
    plan, target = written
    book = load_workbook(target)
    book.worksheets[2].title = "Renamed"
    book.save(target)
    with pytest.raises(WorkbookError, match="'Events' was renamed"):
        read_edits(target, plan)


@pytest.mark.parametrize("damage", ["deleted", "format"])
def test_a_damaged_meta_sheet_is_refused(
    written: tuple[ClassifyPlan, Path], damage: str
) -> None:
    """Without `_meta`, or of another format, it is not a classify workbook."""
    plan, target = written
    book = load_workbook(target)
    if damage == "deleted":
        del book[META_SHEET]
    else:
        book[META_SHEET].cell(MetaKey.FORMAT, 2).value = "99"
    book.save(target)
    with pytest.raises(WorkbookError, match="not a classify workbook"):
        read_edits(target, plan)


def test_the_summary_has_a_free_notes_column(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """Notes are editable on every sheet, the summary included, and ignored."""
    plan, target = written
    book = load_workbook(target)
    summary = book.worksheets[0]
    assert summary.cell(1, 3).value == "Notes"
    assert not summary.cell(2, 3).protection.locked
    assert summary.cell(2, 1).protection.locked
    summary.cell(2, 3).value = "Ask Mum about 2016"
    book.save(target)
    assert not read_edits(target, plan).files
