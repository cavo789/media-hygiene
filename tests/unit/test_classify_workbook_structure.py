"""Every tampering with the workbook's structure is refused, naming what differs."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.sheets import (
    META_SHEET,
    FileColumn,
    Labels,
    MetaKey,
)
from media_hygiene.classify.workbook.structure import language_of
from media_hygiene.classify.workbook.writer import write_workbook
from media_hygiene.constants import Locale
from media_hygiene.errors import WorkbookError
from media_hygiene.i18n import install, using
from tests.unit.test_classify_workbook import make_plan

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook

    from media_hygiene.classify.plan_file import ClassifyPlan


def written(tmp_path: Path) -> tuple[ClassifyPlan, Path]:
    """A plan and its workbook, in `tmp_path`."""
    target, plan = tmp_path / "classify.xlsx", make_plan()
    write_workbook(plan, target)
    return plan, target


def swap_columns(book: Workbook) -> None:
    """Swap the Folder and Name columns of the Files sheet."""
    files = book.worksheets[3]
    for row in files.iter_rows():
        row[2].value, row[3].value = row[3].value, row[2].value


def delete_row(book: Workbook) -> None:
    """Delete the second file."""
    book.worksheets[3].delete_rows(3)


def edit_locked(book: Workbook) -> None:
    """Change a proposed folder."""
    book.worksheets[3].cell(2, FileColumn.PROPOSAL).value = "Elsewhere"


def edit_header(book: Workbook) -> None:
    """Rename a header."""
    book.worksheets[2].cell(1, 2).value = "When"


def duplicate_row(book: Workbook) -> None:
    """Copy the id of the first file onto the second."""
    files = book.worksheets[3]
    files.cell(3, FileColumn.ID).value = files.cell(2, FileColumn.ID).value


def move_sheet(book: Workbook) -> None:
    """Put the Files sheet first."""
    book.move_sheet("Files", offset=-3)


def add_sheet(book: Workbook) -> None:
    """Add a sheet of one's own."""
    book.create_sheet("Mine")


def edit_summary(book: Workbook) -> None:
    """Change a count of the Summary."""
    book.worksheets[0].cell(3, 2).value = 999


def forge_meta(book: Workbook) -> None:
    """Change the fingerprint of `_meta`."""
    book[META_SHEET].cell(MetaKey.FINGERPRINT, 2).value = "0" * 64


@pytest.mark.parametrize(
    ("tamper", "message"),
    [
        (swap_columns, "Files!C1: expected 'Folder', found 'Name'"),
        (delete_row, "Files: the row .* was deleted"),
        (edit_locked, "Files!J2: expected .*, found 'Elsewhere'"),
        (edit_header, "Events!B1: expected 'Dates', found 'When'"),
        (duplicate_row, "Files!A3: the row .* appears twice"),
        (move_sheet, "renamed, moved, added or deleted: Files, Summary"),
        (add_sheet, "renamed, moved, added or deleted: .*Mine"),
        (edit_summary, "Summary!B3: expected"),
        (forge_meta, "do not match"),
    ],
)
def test_each_tampering_is_refused_with_its_cell(
    tmp_path: Path, tamper: object, message: str
) -> None:
    """The first difference refuses the whole workbook, with a tip to undo it."""
    plan, target = written(tmp_path)
    book = load_workbook(target)
    tamper(book)  # type: ignore[operator]
    book.save(target)
    with pytest.raises(WorkbookError, match=message) as raised:
        read_edits(target, plan)
    assert raised.value.tip


def test_a_french_workbook_is_checked_in_french(tmp_path: Path) -> None:
    """Written in French, read under English: its headers are compared in French."""
    plan = make_plan()
    target = tmp_path / "classify.xlsx"
    install(Locale.FR)
    write_workbook(plan, target)
    install(Locale.EN)
    assert not read_edits(target, plan).files
    with using(Locale.FR):
        assert language_of(Labels.current()) is Locale.FR


def test_labels_of_no_known_language_are_refused(tmp_path: Path) -> None:
    """A workbook of another version, whose values no language gives."""
    plan, target = written(tmp_path)
    book = load_workbook(target)
    book[META_SHEET].cell(MetaKey.STAY, 2).value = "(bleibt)"
    book.save(target)
    with pytest.raises(WorkbookError, match="another version"):
        read_edits(target, plan)
