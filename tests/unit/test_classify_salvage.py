"""A workbook `sort` refuses is still salvaged: ids and editable columns by header."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from openpyxl import Workbook, load_workbook

from media_hygiene.classify.carry import carry_over
from media_hygiene.classify.carry_models import EditSheet, LostWhy
from media_hygiene.classify.carry_types import CarrySource, typed_text
from media_hygiene.classify.names import lost_line
from media_hygiene.classify.workbook.edits import CategoryEdit, Stay
from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.salvage import salvage
from media_hygiene.classify.workbook.sheets import (
    META_SHEET,
    CategoryColumn,
    EventColumn,
    FileColumn,
    Labels,
)
from media_hygiene.classify.workbook.writer import write_workbook
from media_hygiene.constants import Locale
from media_hygiene.errors import WorkbookError
from media_hygiene.i18n import using
from tests.unit.test_classify_workbook import make_plan

if TYPE_CHECKING:
    from pathlib import Path

    from media_hygiene.classify.plan_file import ClassifyPlan

_FILES, _EVENTS, _CATEGORIES = 3, 2, 1


def _written(tmp_path: Path, locale: Locale = Locale.EN) -> tuple[ClassifyPlan, Path]:
    """A plan and its workbook, edited on every sheet."""
    plan = make_plan()
    target = tmp_path / "classify.xlsx"
    with using(locale):
        write_workbook(plan, target)
        stay = Labels.current().stay
    book = load_workbook(target)
    files = book.worksheets[_FILES]
    files.cell(2, FileColumn.FINAL).value = "2016/Fair"
    files.cell(3, FileColumn.FINAL).value = stay
    files.cell(4, FileColumn.FINAL).value = "a:b"
    files.cell(5, FileColumn.NOTES).value = "Ask Mum"
    events = book.worksheets[_EVENTS]
    events.cell(2, EventColumn.NAME).value = "Kermesse"
    events.cell(3, EventColumn.CATEGORY).value = stay
    events.cell(3, EventColumn.NOTES).value = "School"
    categories = book.worksheets[_CATEGORIES]
    categories.cell(2, CategoryColumn.RENAME).value = "Beach"
    categories.cell(2, CategoryColumn.CONFIRM).value = "maybe"
    book.save(target)
    return plan, target


@pytest.mark.parametrize("locale", list(Locale))
def test_every_edit_is_read_and_bad_values_set_aside(
    tmp_path: Path, locale: Locale
) -> None:
    """In any language: folders, "stay", names, notes; a bad value is listed."""
    plan, target = _written(tmp_path, locale)
    found = salvage(target)
    assert found.plan_id == plan.plan_id
    assert sorted(found.edits.files.values(), key=str) == ["2016/Fair", Stay.STAY]
    assert [e.name for e in found.edits.events.values()] == ["Kermesse", ""]
    assert list(found.edits.categories.values()) == [CategoryEdit("Beach")]
    assert list(found.notes.files.values()) == ["Ask Mum"]
    assert list(found.notes.events.values()) == ["School"]
    assert {(lost.sheet, lost.value) for lost in found.invalid} == {
        (EditSheet.FILES, "a:b"),
        (EditSheet.CATEGORIES, "maybe"),
    }
    assert not found.empty


def test_a_renamed_sheet_and_a_deleted_column_are_salvaged(tmp_path: Path) -> None:
    """What `sort` refuses: the sheets are found by their headers."""
    plan, target = _written(tmp_path)
    book = load_workbook(target)
    book.worksheets[_EVENTS].title = "My events"
    book.worksheets[_FILES].delete_cols(3)  # the Folder column
    book.move_sheet(book.worksheets[_FILES], offset=-2)
    book.save(target)
    with pytest.raises(WorkbookError):
        read_edits(target, plan)
    found = salvage(target)
    assert "2016/Fair" in found.edits.files.values()
    assert [e.name for e in found.edits.events.values()] == ["Kermesse", ""]


def test_without_meta_or_plan_ids_are_matched_as_they_are(tmp_path: Path) -> None:
    """Row and event ids come from the content: they still match the new plan."""
    plan, target = _written(tmp_path)
    book = load_workbook(target)
    del book[META_SHEET]
    book.save(target)
    found = salvage(target)
    assert not found.plan_id
    carried = carry_over(
        CarrySource("C:\\classify.xlsx", "2026-10-02T14:32", found, None), plan
    )
    assert carried.edits.files == dict(found.edits.files)
    assert set(carried.edits.events) == set(found.edits.events)
    assert carried.count == len(found.edits.files) + 2 + 1
    assert carried.note_count == 2


def test_a_file_that_holds_no_sheet_of_a_workbook_is_refused(tmp_path: Path) -> None:
    """Nothing to salvage: say so, rather than carry nothing silently."""
    target = tmp_path / "other.xlsx"
    Workbook().save(target)
    with pytest.raises(WorkbookError, match="no sheet of a classify workbook"):
        salvage(target)


def test_lost_edits_read_as_one_line(tmp_path: Path) -> None:
    """The console and the report say which sheet, what was typed and why."""
    plan, target = _written(tmp_path)
    found = salvage(target)
    source = CarrySource("C:\\classify.xlsx", "2026-10-02T14:32", found, plan)
    loose = tuple(row.model_copy(update={"event_id": ""}) for row in plan.rows[1:])
    renamed = plan.model_copy(update={"rows": loose, "events": ()})
    carried = carry_over(source, renamed)
    lines = [lost_line(lost) for lost in carried.lost]
    assert any("the value cannot be used" in line for line in lines)
    assert any(line.startswith("Events: ") for line in lines)
    assert {lost.why for lost in carried.lost} >= {LostWhy.GONE, LostWhy.INVALID}
    assert typed_text(Stay.STAY) == "(stay where it is)"
    assert typed_text(CategoryEdit(confirm=True)) == "yes"
