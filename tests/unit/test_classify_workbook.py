"""The classify workbook: structure locked, edits read back, damage refused."""

# classify works on naive local dates, as EXIF writes them.
# ruff: noqa: DTZ001

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from media_hygiene.classify.engine import Scope, classify
from media_hygiene.classify.models import Band, MediaInput
from media_hygiene.classify.plan_build import build_plan
from media_hygiene.classify.workbook.edits import Stay, resolve
from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.sheets import (
    META_SHEET,
    CategoryColumn,
    EventColumn,
    FileColumn,
)
from media_hygiene.classify.workbook.writer import write_workbook
from media_hygiene.classify.worklist import work_list
from media_hygiene.config.classify_settings import ClassifySettings
from media_hygiene.errors import WorkbookError
from media_hygiene.paths.host_paths import HostPathMapper
from media_hygiene.scan.models import VisualFacts

if TYPE_CHECKING:
    from media_hygiene.classify.plan_file import ClassifyPlan

ROOT = Path("/data/c/Photos")
SETTINGS = ClassifySettings()


def shot(relative: str, when: datetime) -> MediaInput:
    """A photo with an EXIF date."""
    visual = VisualFacts(0, 0, 1, 1, 0.0, when.strftime("%Y:%m:%d %H:%M:%S"), "Pixel")
    return MediaInput(ROOT / relative, ROOT, 10, 0, visual, digest=relative)


def make_plan() -> ClassifyPlan:
    """A named folder, a large loose event and a small one."""
    start = datetime(2016, 7, 14, 10)
    files = [shot("2019/Seaside/a.jpg", datetime(2019, 8, 1, 10))]
    files += [
        shot(f"DCIM/big{index}.jpg", start + timedelta(hours=index))
        for index in range(8)
    ]
    files += [
        shot(
            f"DCIM/small{index}.jpg", datetime(2017, 3, 1, 10) + timedelta(hours=index)
        )
        for index in range(5)
    ]
    classification = classify(files, SETTINGS, Scope())
    return build_plan(classification, SETTINGS, HostPathMapper(Path("/data")))


@pytest.fixture(name="written")
def written_workbook(tmp_path: Path) -> tuple[ClassifyPlan, Path]:
    """A plan and its workbook."""
    plan = make_plan()
    target = tmp_path / "classify.xlsx"
    write_workbook(plan, target)
    return plan, target


def test_rows_have_stable_ids_and_host_paths() -> None:
    """Two runs give the same row and event ids; paths are the host's."""
    first, second = make_plan(), make_plan()
    assert [row.id for row in first.rows] == [row.id for row in second.rows]
    assert [event.id for event in first.events] == [e.id for e in second.events]
    assert first.plan_id != second.plan_id
    assert first.rows[0].path.startswith("C:\\Photos\\")


def test_the_work_list_puts_the_largest_undecided_event_first() -> None:
    """Naming the first row covers most of the work."""
    items = work_list(make_plan())
    assert [len(item.rows) for item in items] == [8, 5]
    assert items[0].cumulated == pytest.approx(8 / 13)
    assert items[-1].cumulated == pytest.approx(1.0)


def test_the_structure_is_locked_and_only_edit_cells_are_open(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """Protected sheets, unlocked text cells, drop-downs, a very hidden `_meta`."""
    _plan, target = written
    book = load_workbook(target)
    assert book.security is not None
    assert book.security.lockStructure
    assert book[META_SHEET].sheet_state == "veryHidden"
    files = book.worksheets[3]
    assert files.protection.sheet
    assert not files.protection.autoFilter
    assert files.cell(2, FileColumn.ID).protection.locked
    final = files.cell(2, FileColumn.FINAL)
    assert not final.protection.locked
    assert final.number_format == "@"
    assert files.data_validations.dataValidation


def test_an_untouched_workbook_changes_nothing(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """No edit: every row keeps its proposal."""
    plan, target = written
    edits = read_edits(target, plan)
    assert not edits.files
    assert not edits.events
    assert not edits.categories
    decisions = resolve(plan, edits)
    assert [d.folder for d in decisions] == [row.folder for row in plan.rows]


def test_edits_are_read_back_with_file_over_event(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """Name the largest event, move one of its files elsewhere, leave one in place."""
    plan, target = written
    book = load_workbook(target)
    events, files = book.worksheets[2], book.worksheets[3]
    events.cell(2, EventColumn.NAME).value = "Kermesse"
    ids = [files.cell(row, FileColumn.ID).value for row in range(2, 16)]
    big = [row.id for row in plan.rows if row.name.startswith("big")]
    files.cell(ids.index(big[0]) + 2, FileColumn.FINAL).value = "2016\\Fête"
    files.cell(ids.index(big[1]) + 2, FileColumn.FINAL).value = "(stay where it is)"
    book.save(target)
    edits = read_edits(target, plan)
    assert edits.files == {big[0]: "2016/Fête", big[1]: Stay.STAY}
    decided = {d.row.id: d for d in resolve(plan, edits)}
    assert decided[big[0]].folder == "2016/Fête"
    assert decided[big[1]].folder is None
    assert decided[big[2]].folder == "2016/Kermesse"
    assert decided[big[2]].band is Band.SURE


def test_values_that_cannot_be_used_name_their_cell(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """A folder climbing out of the target, a confirmation that is not yes/no."""
    plan, target = written
    book = load_workbook(target)
    book.worksheets[3].cell(3, FileColumn.FINAL).value = "..\\Windows"
    book.worksheets[1].cell(2, CategoryColumn.CONFIRM).value = "maybe"
    book.save(target)
    with pytest.raises(WorkbookError) as raised:
        read_edits(target, plan)
    assert "Files!K3" in raised.value.message
    assert "Categories!E2" in raised.value.message


def test_a_workbook_of_another_run_is_refused(
    written: tuple[ClassifyPlan, Path],
) -> None:
    """The plan id of `_meta` must match `plan.json`."""
    _plan, target = written
    with pytest.raises(WorkbookError, match="another classify run"):
        read_edits(target, make_plan())


def test_a_changed_locked_cell_is_refused(written: tuple[ClassifyPlan, Path]) -> None:
    """Rows are matched by id: a forged id is not a row of the plan."""
    plan, target = written
    book = load_workbook(target)
    book.worksheets[3].cell(2, FileColumn.ID).value = "forged"
    book.save(target)
    with pytest.raises(WorkbookError, match="Files!A2: 'forged' is not a row"):
        read_edits(target, plan)
