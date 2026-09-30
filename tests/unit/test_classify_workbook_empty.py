"""A workbook sheet without rows: headers alone, read back (TODO 0047)."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from media_hygiene.classify.models import Band
from media_hygiene.classify.workbook.reader import read_edits
from media_hygiene.classify.workbook.salvage import salvage
from media_hygiene.classify.workbook.writer import write_workbook
from tests.unit.test_classify_workbook import make_plan

if TYPE_CHECKING:
    from pathlib import Path

    from media_hygiene.classify.plan_file import ClassifyPlan

_CATEGORIES, _EVENTS, _FILES = 1, 2, 3


def _no_events() -> ClassifyPlan:
    """No event reaches `min_event_size`: the Events sheet is empty."""
    return make_plan().model_copy(update={"events": ()})


def _no_categories() -> ClassifyPlan:
    """Every file stays where it is: the Categories sheet is empty."""
    plan = make_plan()
    rows = tuple(
        row.model_copy(update={"values": None, "band": Band.STAY, "folder": None})
        for row in plan.rows
    )
    return plan.model_copy(update={"rows": rows, "events": ()})


def _no_rows() -> ClassifyPlan:
    """A mounted folder without media: every sheet but the Summary is empty."""
    return make_plan().model_copy(update={"rows": (), "events": ()})


@pytest.mark.parametrize(
    ("plan", "empty"),
    [
        (_no_events(), (_EVENTS,)),
        (_no_categories(), (_CATEGORIES, _EVENTS)),
        (_no_rows(), (_CATEGORIES, _EVENTS, _FILES)),
    ],
    ids=["no-event", "no-category", "no-file"],
)
def test_an_empty_sheet_keeps_its_headers_and_reads_back(
    plan: ClassifyPlan, empty: tuple[int, ...], tmp_path: Path
) -> None:
    """No crash: headers alone, no drop-down, the filter on the header row."""
    target = tmp_path / "classify.xlsx"
    write_workbook(plan, target)
    book = load_workbook(target)
    for index in empty:
        sheet = book.worksheets[index]
        assert sheet.max_row == 1, sheet.title
        assert sheet.cell(1, 1).value
        assert not sheet.data_validations.dataValidation
        assert sheet.auto_filter.ref is not None
        assert sheet.auto_filter.ref.endswith("1")
    edits = read_edits(target, plan)
    assert not edits.files
    assert not edits.events
    assert not edits.categories
    assert salvage(target).plan_id == plan.plan_id


def test_the_summary_of_a_plan_without_files_says_zero(tmp_path: Path) -> None:
    """`0 of 0` would divide by zero: the progress line says 0."""
    target = tmp_path / "classify.xlsx"
    write_workbook(_no_rows(), target)
    summary = load_workbook(target).worksheets[0]
    assert summary.cell(2, 2).value == 0
