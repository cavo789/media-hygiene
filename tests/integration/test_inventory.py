"""`inventory`: the index exported to Excel or CSV, without reading any file again."""

from __future__ import annotations

import builtins
import csv
import shutil
from datetime import datetime
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from media_hygiene.constants import InventoryFormat, Locale
from media_hygiene.i18n import install
from media_hygiene.scan.progress import NullProgress
from media_hygiene.services.audit import AuditService
from media_hygiene.services.inventory import export_inventory
from tests.support.runtime import make_runtime
from tests.support.scenes import Effect, Shot, write_shot

if TYPE_CHECKING:
    from pathlib import Path

    import pytest
    from openpyxl.worksheet.worksheet import Worksheet

    from media_hygiene.paths.locations import Locations
    from media_hygiene.services.runtime import Runtime
    from tests.support.media import MediaFactory


def audited(locations: Locations, media: MediaFactory) -> Runtime:
    """A dated photo and its copy, a blurred undated photo, a video; then an audit."""
    data = locations.data_dir
    dated = write_shot(data / "c/Photos/a.jpg", Shot(1, taken_at="2021:06:05 12:00:00"))
    media.copy(dated, "d/Backup/a.jpg")
    write_shot(data / "c/Photos/b.jpg", Shot(2, camera=False, effect=Effect.BLURRED))
    media.video("c/Videos/v.mp4")
    runtime = make_runtime(locations)
    AuditService(runtime, NullProgress()).run()
    return runtime


def rows(sheet: Worksheet) -> list[tuple[object, ...]]:
    """Every row of a sheet, as values."""
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_the_workbook_lists_every_file_with_typed_cells(
    locations: Locations, media: MediaFactory
) -> None:
    """One row per file, host paths, typed cells, frozen header, filter, groups."""
    result = export_inventory(audited(locations, media), InventoryFormat.XLSX)
    assert result.target.name == "inventory.xlsx"
    assert result.target.parent.name.endswith("-inventory")
    book = load_workbook(result.target)
    assert book.sheetnames == ["Files", "Summary"]
    files = book["Files"]
    assert files.freeze_panes == "A2"
    assert files.auto_filter.ref == f"A1:AP{files.max_row}"
    header, *lines = rows(files)
    by_path = {line[0]: dict(zip(header, line, strict=True)) for line in lines}
    assert set(by_path) == {
        "C:\\Photos\\a.jpg",
        "C:\\Photos\\b.jpg",
        "C:\\Videos\\v.mp4",
        "D:\\Backup\\a.jpg",
    }
    photo = by_path["C:\\Photos\\a.jpg"]
    assert photo["Date taken"] == datetime(2021, 6, 5, 12)  # noqa: DTZ001 - naive
    assert (photo["Year"], photo["Date from"], photo["Kind"]) == (2021, "EXIF", "photo")
    assert photo["Camera"] == "Canon EOS 80D"
    assert isinstance(photo["Size (bytes)"], int)
    assert isinstance(photo["Modified (UTC)"], datetime)
    copy = by_path["D:\\Backup\\a.jpg"]
    assert (photo["Duplicate group"], photo["Copies"]) == (1, 2)
    assert (copy["Duplicate group"], copy["SHA-256"]) == (1, photo["SHA-256"])
    blurred = by_path["C:\\Photos\\b.jpg"]
    assert "blurry" in str(blurred["Quality"])
    assert blurred["Date taken"] is None
    assert blurred["Duplicate group"] is None
    video = by_path["C:\\Videos\\v.mp4"]
    assert (video["Kind"], video["Date from"]) == ("video", "video tags")
    assert video["Duration (s)"] == 2  # a number, not text
    assert all(line["Integrity"] == "healthy" for line in by_path.values())


def test_the_summary_says_how_fresh_each_root_is(
    locations: Locations, media: MediaFactory
) -> None:
    """The Summary sheet dates the last complete audit of every root, then counts."""
    result = export_inventory(audited(locations, media), InventoryFormat.XLSX)
    summary = rows(load_workbook(result.target)["Summary"])
    assert summary[0] == ("Folder", "Last complete audit (UTC)", None)
    dated = [line for line in summary if isinstance(line[1], datetime)]
    assert dated
    assert len(dated) == len(result.roots)
    assert ("Kind", "photo", 3) in summary
    assert ("Duplicates", "groups of identical files", 1) in summary
    assert result.summary.files == len(rows(load_workbook(result.target)["Files"])) - 1


def test_the_csv_holds_the_same_rows(locations: Locations, media: MediaFactory) -> None:
    """`--format csv` writes the Files sheet, as `plan.csv` does (BOM, `,`)."""
    runtime = audited(locations, media)
    book = load_workbook(export_inventory(runtime, InventoryFormat.XLSX).target)
    target = export_inventory(runtime, InventoryFormat.CSV).target
    assert target.name == "inventory.csv"
    assert target.read_bytes().startswith(b"\xef\xbb\xbf")
    with target.open(encoding="utf-8-sig", newline="") as stream:
        lines = list(csv.reader(stream))
    sheet = rows(book["Files"])
    # Path, folder, name, kind and size: the same rows in the same order.
    assert [line[:5] for line in lines] == [
        ["" if cell is None else str(cell) for cell in row[:5]] for row in sheet
    ]


def test_no_file_of_data_is_opened(
    locations: Locations, media: MediaFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The export reads the index only: the files may even be gone."""
    runtime = audited(locations, media)
    shutil.rmtree(locations.data_dir)
    locations.data_dir.mkdir()
    opened: list[str] = []
    real_open = builtins.open

    def counting(file: object, *args: object, **kwargs: object) -> object:
        opened.append(str(file))
        return real_open(file, *args, **kwargs)  # type: ignore[call-overload]

    monkeypatch.setattr(builtins, "open", counting)
    result = export_inventory(runtime, InventoryFormat.XLSX)
    assert result.summary.files == 4
    assert opened  # the workbook itself was written
    assert not [path for path in opened if path.startswith(str(locations.data_dir))]


def test_headers_and_sheets_are_translated(
    locations: Locations, media: MediaFactory
) -> None:
    """In French, the sheets and headers are French; the CSV uses `;`."""
    runtime = audited(locations, media)
    install(Locale.FR)
    book = load_workbook(export_inventory(runtime, InventoryFormat.XLSX).target)
    assert book.sheetnames == ["Fichiers", "Résumé"]
    assert rows(book["Fichiers"])[0][:2] == ("Fichier", "Dossier")
    target: Path = export_inventory(runtime, InventoryFormat.CSV).target
    assert target.read_text(encoding="utf-8-sig").startswith("Fichier;Dossier;")
