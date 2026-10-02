"""`sort` applies the edited workbook, proves nothing was lost; `undo` reverses it."""

from __future__ import annotations

import json
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from media_hygiene.actions.journal import journal_file, read_journal
from media_hygiene.actions.kinds import ActionKind, Phase
from media_hygiene.actions.runs import list_run_ids
from media_hygiene.classify.models import Band, SortReason
from media_hygiene.classify.workbook.sheets import FileColumn
from media_hygiene.constants import Locale, MediaKind
from media_hygiene.errors import WorkbookError
from media_hygiene.i18n import install
from media_hygiene.index.repository import FactsRepository
from media_hygiene.scan.models import MediaFile
from media_hygiene.scan.progress import NullProgress
from media_hygiene.services.audit import AuditService
from media_hygiene.services.classify import ClassifyService
from tests.support.docs.library import build_library as build_docs_library
from tests.support.docs.names import Locale as DocsLocale
from tests.support.runtime import make_runtime
from tests.support.scenes import Shot, write_shot
from tests.support.sorting import (
    EVENT_NAME,
    MIXED,
    PARTY,
    WEDDING,
    build_library,
    classify,
    folders,
    name_first_event,
    snapshot,
    sort,
    undo,
)

if TYPE_CHECKING:
    from typer.testing import CliRunner

    from media_hygiene.paths.locations import Locations


def test_classify_edit_sort_then_undo_gives_the_tree_back(locations: Locations) -> None:
    """Moves as edited, nothing lost, then byte-identical after undo, folders back."""
    data = locations.data_dir
    build_library(data)
    runtime = make_runtime(locations)
    AuditService(runtime, NullProgress()).run()
    start, tree = snapshot(data), folders(data)
    workbook = classify(runtime)
    name_first_event(workbook)
    result = sort(runtime)
    assert result.manifest.intact, result.manifest
    assert not result.moves.outcome.failed
    assert (data / "c/2016" / EVENT_NAME / "IMG_0000.jpg").is_file()
    assert (data / "c/2018/Mariage/DSC_0000.xmp").is_file()  # the sidecar followed
    assert not (data / PARTY).exists()  # Thumbs.db quarantined, folder removed
    assert not (data / WEDDING).exists()  # desktop.ini too
    assert (data / MIXED / "notes.txt").is_file()  # a folder holding more stays
    assert result.folders.removed >= 2
    assert isinstance(result.manifest_file, Path)
    manifest = json.loads(result.manifest_file.read_text("utf-8"))
    assert manifest["before"] == manifest["after"]
    outcome = undo(runtime, result.run_id)
    assert not outcome.failed, outcome
    assert not outcome.skipped, outcome
    assert snapshot(data) == start
    assert folders(data) == tree


def test_the_index_follows_the_moves_and_undo(locations: Locations) -> None:
    """A moved file keeps its digest: the next audit hashes nothing again."""
    build_library(locations.data_dir)
    runtime = make_runtime(locations)
    AuditService(runtime, NullProgress()).run()
    name_first_event(classify(runtime))
    result = sort(runtime)
    moved = next(e for e in result.moves.moved if e.path.endswith("IMG_0000.jpg"))
    stat = (moved.size, moved.mtime_ns, MediaKind.IMAGE)
    with FactsRepository.open(runtime.index_file) as index:
        assert index.get(MediaFile(Path(moved.target or ""), *stat)).integrity
        assert not index.get(MediaFile(Path(moved.path), *stat)).integrity
    undo(runtime, result.run_id)
    with FactsRepository.open(runtime.index_file) as index:
        assert index.get(MediaFile(Path(moved.path), *stat)).integrity


def test_classify_right_after_sort_proposes_nothing_to_move(
    locations: Locations,
) -> None:
    """In place: every file already sits where the plan puts it."""
    build_library(locations.data_dir)
    runtime = make_runtime(locations)
    name_first_event(classify(runtime))
    sort(runtime)
    proposals = ClassifyService(runtime, NullProgress()).run().classification.proposals
    moving = [p.file.path for p in proposals if not p.in_place]
    assert not moving


def test_an_unedited_sort_then_classify_proposes_nothing_to_move(
    cli: CliRunner, locations: Locations
) -> None:
    """Every band, "to check" included: the guesses sorted stay where they are."""
    assert cli is not None  # the demo tree, in the data folder
    runtime = make_runtime(locations)
    classify(runtime)
    unsure = [
        p
        for p in ClassifyService(runtime, NullProgress()).run().classification.proposals
        if p.band is Band.UNSURE and not p.in_place
    ]
    assert unsure  # the demo tree has guesses to check
    assert not sort(runtime).moves.outcome.failed
    proposals = ClassifyService(runtime, NullProgress()).run().classification.proposals
    assert not [p.file.path for p in proposals if not p.in_place]
    kept = [p for p in proposals if p.reason is SortReason.PREVIOUS_GUESS]
    assert len(kept) == len(unsure)


def test_the_journal_names_the_plan_and_each_row(locations: Locations) -> None:
    """A resumed run finds the rows of its plan in the journal."""
    build_library(locations.data_dir)
    runtime = make_runtime(locations)
    name_first_event(classify(runtime))
    result = sort(runtime)
    entries = read_journal(journal_file(locations.journal_dir, result.run_id))
    moves = [
        e for e in entries if e.action is ActionKind.MOVE and e.phase is Phase.SORT
    ]
    assert {e.plan for e in moves} == {result.manifest.plan_id}
    assert all(e.row for e in moves if not e.path.endswith(".xmp"))


def type_folders(workbook: Path, typed: dict[str, str]) -> None:
    """Type a final folder on the Files rows of these file names, as in Excel."""
    book = load_workbook(workbook)
    for row in book.worksheets[3].iter_rows(min_row=2):
        names = {str(cell.value).rsplit("\\", 1)[-1] for cell in row if cell.value}
        for name in names & typed.keys():
            row[FileColumn.FINAL - 1].value = typed[name]
    book.save(workbook)


def test_a_file_edit_on_a_twin_decides_for_both(locations: Locations) -> None:
    """IMG_0000.jpg edited, not the leading IMG_0000.jpeg: both go where it says."""
    data = locations.data_dir
    build_library(data)
    write_shot(data / PARTY / "IMG_0000.jpeg", Shot(90, taken_at="2016:07:14 10:00:00"))
    runtime = make_runtime(locations)
    workbook = classify(runtime)
    start = snapshot(data)
    type_folders(workbook, {"IMG_0000.jpg": "Vacances", "IMG_0001.jpg": "Plage"})
    assert sort(runtime).manifest.intact
    assert (data / "c/Vacances/IMG_0000.jpg").is_file()
    assert (data / "c/Vacances/IMG_0000.jpeg").is_file()
    type_folders(workbook, {"IMG_0000.jpeg": "Mer"})  # now two cells disagree
    for run_id in list_run_ids(locations.journal_dir):
        undo(runtime, run_id)
    assert snapshot(data) == start
    with pytest.raises(WorkbookError) as refused:
        sort(runtime)
    message = " ".join(refused.value.message.split())
    assert "(IMG_0000.jpg)" in message
    assert "(IMG_0000.jpeg)" in message
    assert snapshot(data) == start  # nothing moved


@pytest.mark.parametrize("locale", list(DocsLocale))
def test_the_documentation_library_without_clean_settles_after_one_sort(
    locations: Locations, locale: DocsLocale
) -> None:
    """Copies left in place: "to sort" files named after their event stay put."""
    install(Locale(locale.value))
    build_docs_library(locale, locations.data_dir)
    runtime = make_runtime(locations)
    classify(runtime)
    assert not sort(runtime).moves.outcome.failed
    proposals = ClassifyService(runtime, NullProgress()).run().classification.proposals
    assert not [p.file.path for p in proposals if not p.in_place]
