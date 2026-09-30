"""A library too small for an event: a workbook `sort` accepts (TODO 0047)."""

from __future__ import annotations

import shutil
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from media_hygiene.scan.progress import NullProgress
from media_hygiene.services.audit import AuditService
from tests.support.cli import run
from tests.support.runtime import make_runtime
from tests.support.scenes import Shot, write_shot
from tests.support.sorting import PHOTOS, classify, snapshot, sort

if TYPE_CHECKING:
    from typer.testing import CliRunner

    from media_hygiene.paths.locations import Locations

_EVENTS = 2


def test_two_photos_give_a_workbook_that_sort_accepts(locations: Locations) -> None:
    """No event reaches its minimum size: the Events sheet has its headers alone."""
    data = locations.data_dir
    for index in range(2):
        shot = Shot(10 + index, taken_at=f"2016:07:14 1{index}:00:00")
        write_shot(data / PHOTOS / "2016" / f"IMG_{index:04d}.jpg", shot)
    runtime = make_runtime(locations)
    AuditService(runtime, NullProgress()).run()
    workbook = classify(runtime)
    assert load_workbook(workbook).worksheets[_EVENTS].max_row == 1
    before = snapshot(data)
    result = sort(runtime)
    assert result.manifest.intact, result.manifest
    assert not result.moves.outcome.failed
    assert len(snapshot(data)) == len(before)


def test_the_demo_tree_after_clean_gets_its_workbook(
    cli: CliRunner, locations: Locations
) -> None:
    """The case 0047 was seen on: no event left once the copies are gone."""
    assert run(cli, "clean", "--yes").exit_code == 0
    result = run(cli, "classify")
    assert result.exit_code == 0, result.output
    (workbook,) = locations.reports_dir.glob("*-classify/classify.xlsx")
    assert load_workbook(workbook).worksheets[_EVENTS].max_row == 1
    again = run(cli, "classify")  # carries the edits of that workbook over
    assert again.exit_code == 0, again.output
    assert run(cli, "sort", "--yes").exit_code == 0


def test_a_folder_without_media_writes_no_workbook(
    cli: CliRunner, locations: Locations
) -> None:
    """Nothing to propose: said on screen, no plan, workbook nor report."""
    data = locations.data_dir
    shutil.rmtree(data)
    (data / PHOTOS).mkdir(parents=True)
    (data / PHOTOS / "notes.txt").write_text("no photo here", encoding="utf-8")
    result = run(cli, "classify")
    assert result.exit_code == 0, result.output
    assert "No photo nor video to classify" in result.output
    assert not list(locations.reports_dir.glob("*-classify"))
