r"""A library whose events split and merge with `merge_gap_hours`, and workbook edits.

`C:\\Photos\\DCIM` holds a morning (5 shots) and an afternoon (6 shots) of one day:
one event with the default gaps, two with `merge_gap_hours = 1`. `C:\\Photos\\Mariage`
holds three shots of a wedding: the category `Mariage`.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Final

from openpyxl import load_workbook

from media_hygiene.classify.plan_file import ClassifyPlan
from media_hygiene.classify.workbook.sheets import (
    CategoryColumn,
    EventColumn,
    FileColumn,
)
from media_hygiene.scan.progress import NullProgress
from media_hygiene.services.carry import CarryRequest, carry_source
from media_hygiene.services.classify import ClassifyService
from media_hygiene.services.classify_output import write_classify_output
from tests.support.scenes import Shot, write_shot

if TYPE_CHECKING:
    from pathlib import Path

    from media_hygiene.services.classify_output import ClassifyOutput
    from media_hygiene.services.runtime import Runtime

DCIM: Final = "c/Photos/DCIM"
WEDDING: Final = "c/Photos/Mariage"
MORNING: Final = 5
AFTERNOON: Final = 6
SPLIT: Final[dict[str, dict[str, object]]] = {"classify": {"merge_gap_hours": 1}}
_NAME_COLUMN: Final = 4  # the Files sheet: the file name
_FILES, _EVENTS, _CATEGORIES = 3, 2, 1  # sheet positions


def build_day(data_dir: Path) -> None:
    """Write the library below the data folder.

    Args:
        data_dir: The data mount point.
    """
    for index in range(MORNING):
        shot = Shot(10 + index, taken_at=f"2017:05:06 08:{index}0:00")
        write_shot(data_dir / DCIM / f"A{index}.jpg", shot)
    for index in range(AFTERNOON):
        shot = Shot(20 + index, taken_at=f"2017:05:06 16:{index}0:00")
        write_shot(data_dir / DCIM / f"B{index}.jpg", shot)
    for index in range(3):
        shot = Shot(40 + index, taken_at=f"2018:06:02 1{index}:30:00")
        write_shot(data_dir / WEDDING / f"M{index}.jpg", shot)


def run_classify(
    runtime: Runtime, request: CarryRequest | None = None
) -> ClassifyOutput:
    """Run `classify` as the command does: carry, propose, write.

    Args:
        runtime: A runtime whose reports mount persists.
        request: `--carry-over` / `--no-carry-over`; the default carries the latest.

    Returns:
        What was written.
    """
    source = carry_source(runtime, request or CarryRequest())
    result = ClassifyService(runtime, NullProgress()).run()
    written = write_classify_output(runtime, result, source)
    assert written is not None
    return written


def plan_of(written: ClassifyOutput) -> ClassifyPlan:
    """Read the plan a run wrote.

    Args:
        written: The run.

    Returns:
        Its plan.
    """
    text = (written.folder / "plan.json").read_text(encoding="utf-8")
    return ClassifyPlan.model_validate_json(text)


def edit(workbook: Path, cells: dict[tuple[int, str], dict[int, str]]) -> None:
    """Type values in a workbook, as a user does in Excel.

    Args:
        workbook: The workbook.
        cells: (sheet position, key) → column → value. The key of a Files row is its
            file name; of an Events row, its position (`"2"`); of a Categories row,
            its category.
    """
    book = load_workbook(workbook)
    for (position, key), values in cells.items():
        sheet = book.worksheets[position]
        column = _NAME_COLUMN if position == _FILES else 1
        row = next(
            (
                r
                for r in range(2, sheet.max_row + 1)
                if sheet.cell(r, column).value == key
            ),
            int(key) if key.isdigit() else 0,
        )
        for number, value in values.items():
            sheet.cell(row, number).value = value
    book.save(workbook)


def final(name: str, folder: str) -> dict[tuple[int, str], dict[int, str]]:
    """The Final folder of one file.

    Args:
        name: The file name.
        folder: The folder typed.

    Returns:
        The cell to type.
    """
    return {(_FILES, name): {FileColumn.FINAL: folder}}


def event_name(position: int, name: str) -> dict[tuple[int, str], dict[int, str]]:
    """The name of the event on a row of the Events sheet.

    Args:
        position: The row (2 is the first event).
        name: The name typed.

    Returns:
        The cell to type.
    """
    return {(_EVENTS, str(position)): {EventColumn.NAME: name}}


def rename(category: str, new: str) -> dict[tuple[int, str], dict[int, str]]:
    """A new name for a category.

    Args:
        category: The category proposed.
        new: Its new name.

    Returns:
        The cell to type.
    """
    return {(_CATEGORIES, category): {CategoryColumn.RENAME: new}}
