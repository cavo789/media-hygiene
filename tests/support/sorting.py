r"""A small library to sort, and the steps of `classify` → edit → `sort` → `undo`.

`C:\\Photos` holds a loose July 2016 party (six shots in a date folder, with a
`Thumbs.db`), a wedding folder (three shots, a sidecar and a `desktop.ini`) and a folder
holding a text file besides its photo: it must stay.
"""

from __future__ import annotations

import hashlib
from typing import TYPE_CHECKING, Final

from openpyxl import load_workbook

from media_hygiene.classify.workbook.sheets import EventColumn
from media_hygiene.scan.progress import NullProgress
from media_hygiene.services.classify import ClassifyService
from media_hygiene.services.classify_output import write_classify_output
from media_hygiene.services.sort import SortService
from media_hygiene.services.sort_inputs import find_workbook, load_inputs
from media_hygiene.services.undo import undo_run
from tests.support.scenes import Shot, write_shot

if TYPE_CHECKING:
    from collections.abc import Callable
    from pathlib import Path

    from media_hygiene.actions.outcome import Outcome
    from media_hygiene.services.runtime import Runtime
    from media_hygiene.services.sort import SortResult

PHOTOS: Final = "c/Photos"
PARTY: Final = f"{PHOTOS}/2016/Juillet 2016"
WEDDING: Final = f"{PHOTOS}/Mariage"
MIXED: Final = f"{PHOTOS}/Divers"
PARTY_SHOTS: Final = 6
EVENT_NAME: Final = "Kermesse"

type Snapshot = dict[str, tuple[str, int]]


def build_library(data_dir: Path) -> None:
    """Write the library below the data folder.

    Args:
        data_dir: The data mount point.
    """
    for index in range(PARTY_SHOTS):
        shot = Shot(10 + index, taken_at=f"2016:07:14 1{index}:00:00")
        write_shot(data_dir / PARTY / f"IMG_{index:04d}.jpg", shot)
    (data_dir / PARTY / "Thumbs.db").write_bytes(b"thumbnails")
    for index in range(3):
        shot = Shot(30 + index, taken_at=f"2018:06:02 1{index}:30:00")
        write_shot(data_dir / WEDDING / f"DSC_{index:04d}.jpg", shot)
    (data_dir / WEDDING / "DSC_0000.xmp").write_text("<x:xmpmeta/>", encoding="utf-8")
    (data_dir / WEDDING / "desktop.ini").write_text("[.ShellClassInfo]", "utf-8")
    write_shot(
        data_dir / MIXED / "IMG_9000.jpg", Shot(50, taken_at="2019:03:01 09:00:00")
    )
    (data_dir / MIXED / "notes.txt").write_text("to keep", encoding="utf-8")


def snapshot(root: Path) -> Snapshot:
    """Every file below a folder: its SHA-256 and modification time.

    Args:
        root: The folder.

    Returns:
        Relative path → (digest, mtime).
    """
    return {
        path.relative_to(root).as_posix(): (
            hashlib.sha256(path.read_bytes()).hexdigest(),
            path.stat().st_mtime_ns,
        )
        for path in root.rglob("*")
        if path.is_file()
    }


def folders(root: Path) -> set[str]:
    """Every folder below a folder.

    Args:
        root: The folder.

    Returns:
        Their relative paths.
    """
    return {
        path.relative_to(root).as_posix() for path in root.rglob("*") if path.is_dir()
    }


def classify(runtime: Runtime) -> Path:
    """Run `classify` and write its plan and workbook.

    Args:
        runtime: A runtime whose reports mount persists.

    Returns:
        The workbook.
    """
    result = ClassifyService(runtime, NullProgress()).run()
    folder = write_classify_output(runtime, result)
    assert folder is not None
    return folder / "classify.xlsx"


def name_first_event(workbook: Path, name: str = EVENT_NAME) -> None:
    """Name the first event of the work list, as a user does in Excel.

    Args:
        workbook: The workbook.
        name: The name typed.
    """
    book = load_workbook(workbook)
    book.worksheets[2].cell(2, EventColumn.NAME).value = name
    book.save(workbook)


def sort(
    runtime: Runtime,
    workbook: str | None = None,
    stop: Callable[[], bool] = lambda: False,
) -> SortResult:
    """Run `sort` through its service, as the command does.

    Args:
        runtime: The runtime.
        workbook: The workbook, or None for the latest.
        stop: Tells when to stop.

    Returns:
        The result.
    """
    service = SortService(runtime, NullProgress())
    service.ensure_ready()
    inputs = load_inputs(runtime, find_workbook(runtime, workbook))
    return service.execute(service.prepare(inputs), stop)


def undo(runtime: Runtime, run_id: str) -> Outcome:
    """Undo a run.

    Args:
        runtime: The runtime.
        run_id: The run.

    Returns:
        What was restored.
    """
    return undo_run(runtime, run_id, NullProgress())
