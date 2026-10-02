"""Write `inventory.xlsx` (sheets Files and Summary) or `inventory.csv` (files only).

openpyxl's write-only mode streams the rows: 70,000 files stay fast and small in
memory. Cells are typed (numbers, dates) so that sorting and filtering work in any
regional settings; the CSV writes them as text, like `plan.csv`.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from typing import TYPE_CHECKING, Final, cast

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from media_hygiene.classify.workbook.cells import header_cell, text_safe_cell
from media_hygiene.i18n import _
from media_hygiene.report.csv_export import (
    CSV_DATE_FORMAT,
    CSV_ENCODING,
    list_separator,
    spreadsheet_text,
)
from media_hygiene.report.inventory_columns import DATE_FORMAT, columns
from media_hygiene.report.inventory_summary import InventorySummary

if TYPE_CHECKING:
    from collections.abc import Iterable, Sequence
    from pathlib import Path

    from openpyxl.cell.cell import Cell
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet

    from media_hygiene.classify.workbook.cells import WriteOnly
    from media_hygiene.report.inventory_columns import Cell as Value
    from media_hygiene.report.inventory_columns import Column
    from media_hygiene.report.inventory_entries import Entry

_DECIMAL_POINT: Final = "."
_SUMMARY_WIDTHS: Final = (60, 30, 12)


@dataclass(frozen=True, slots=True)
class InventorySource:
    """What an export writes: the files, and when each root was last walked."""

    entries: Iterable[Entry]
    roots: Sequence[tuple[str, datetime]]


def write_workbook(target: Path, source: InventorySource) -> InventorySummary:
    """Write the Files and Summary sheets.

    Args:
        target: The `.xlsx` file to write.
        source: The files and the roots.

    Returns:
        The counts, to show in the console too.
    """
    book = Workbook(write_only=True)
    spec = columns()
    created: WriteOnlyWorksheet = book.create_sheet(_("Files"))
    sheet = cast("WriteOnly", created)
    sheet.freeze_panes = "A2"
    for index, column in enumerate(spec, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.append([header_cell(created, column.header) for column in spec])
    summary, last = InventorySummary(), 1
    for entry in source.entries:
        sheet.append([_cell(created, column, entry) for column in spec])
        summary.add(entry)
        last += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(spec))}{last}"
    _write_summary(book, summary, source.roots)
    book.save(target)
    return summary


def write_csv(target: Path, source: InventorySource) -> InventorySummary:
    """Write the Files sheet as a CSV file, readable by Excel in the interface language.

    Args:
        target: The `.csv` file to write.
        source: The files (the roots are not written).

    Returns:
        The counts, to show in the console.
    """
    spec, summary = columns(), InventorySummary()
    separator = list_separator()
    with target.open("w", encoding=CSV_ENCODING, newline="") as stream:
        writer = csv.writer(stream, delimiter=separator)
        writer.writerow(column.header for column in spec)
        for entry in source.entries:
            writer.writerow(_text(column.value(entry), separator) for column in spec)
            summary.add(entry)
    return summary


def _cell(sheet: WriteOnlyWorksheet, column: Column, entry: Entry) -> Cell:
    """One typed cell of the Files sheet.

    Args:
        sheet: The sheet.
        column: Its column.
        entry: The file of the row.

    Returns:
        The cell, with the number format of its column; text stays text.
    """
    value = column.value(entry)
    written = text_safe_cell(sheet, value)
    if column.number_format is not None:
        written.number_format = column.number_format
    return written


def _text(value: Value, separator: str) -> str:
    """Write a value as Excel reads it with this list separator.

    Args:
        value: The typed value.
        separator: `;` (decimal comma) or `,` (decimal point).

    Returns:
        The text; a name that starts like a formula keeps a leading apostrophe.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(CSV_DATE_FORMAT)
    if isinstance(value, float) and separator != ",":
        return repr(value).replace(_DECIMAL_POINT, ",")
    if isinstance(value, str):
        return spreadsheet_text(value)
    return str(value)


def _write_summary(
    book: Workbook,
    summary: InventorySummary,
    roots: Sequence[tuple[str, datetime]],
) -> None:
    """Add the Summary sheet: how fresh each root is, then the counts.

    Args:
        book: The workbook, write-only.
        summary: The counts.
        roots: Each walked root on the host, and when it was last walked.
    """
    created: WriteOnlyWorksheet = book.create_sheet(_("Summary"))
    sheet = cast("WriteOnly", created)
    for index, width in enumerate(_SUMMARY_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    headers = (_("Folder"), _("Last complete audit (UTC)"))
    sheet.append([header_cell(created, header) for header in headers])
    for folder, walked_at in roots:
        when = text_safe_cell(created, walked_at.replace(tzinfo=None))
        when.number_format = DATE_FORMAT
        sheet.append([text_safe_cell(created, folder), when])
    if not roots:
        sheet.append([_("No audit has walked a folder completely yet.")])
    sheet.append([])
    counts = (_("Topic"), _("Value"), _("Files"))
    sheet.append([header_cell(created, header) for header in counts])
    for row in summary.rows():
        sheet.append([text_safe_cell(created, value) for value in row])
