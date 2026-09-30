"""Write `classify.xlsx`: every sheet protected, only the yellow cells editable.

The structure is locked (no sheet renamed, moved, deleted or added); filtering stays
allowed. Editable cells are unlocked and use the Text format, so `2021` stays a folder
name. Protection is a guard-rail against slips, not security: the reader checks
everything. openpyxl's write-only mode keeps both protection and validation, and
writes 70,000 rows in about seven seconds (read back in about four).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Final, cast

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation

from media_hygiene.classify.workbook import rows as content
from media_hygiene.classify.workbook.cells import (
    editable_cell,
    header_cell,
    locked_cell,
    padded,
)
from media_hygiene.classify.workbook.sheets import (
    META_LIST_COLUMN,
    META_SHEET,
    CategoryColumn,
    EventColumn,
    FileColumn,
    Labels,
    fingerprint,
)
from media_hygiene.classify.workbook.summary import summary_headers, summary_rows

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet._write_only import WriteOnlyWorksheet

    from media_hygiene.classify.plan_file import ClassifyPlan
    from media_hygiene.classify.workbook.cells import WriteOnly
    from media_hygiene.classify.workbook.rows import Row

_PERCENT: Final = "0%"
_WIDTH: Final = 18
_WIDE: Final = 40
_SUMMARY_WIDTHS: Final = (60, 24, _WIDE)
_SUMMARY_NOTES: Final = 3


@dataclass(frozen=True, slots=True)
class _Spec:
    """One protected sheet: its headers, its editable columns, its drop-downs."""

    title: str
    headers: Row
    editable: frozenset[int]
    lists: dict[int, str] = field(default_factory=dict)  # column → list formula
    formats: dict[int, str] = field(default_factory=dict)  # column → number format
    widths: tuple[int, ...] = ()  # the first columns' widths; the others: _WIDTH
    filtered: bool = True  # an auto-filter on the header row


def write_workbook(plan: ClassifyPlan, target: Path) -> None:
    """Write the workbook of a plan.

    Args:
        plan: The plan.
        target: The `.xlsx` file to write.
    """
    labels = Labels.current()
    book = Workbook(write_only=True)
    book.security = WorkbookProtection(lockStructure=True)
    keys: list[str] = list(plan.categories())
    for spec, rows in _sheets(plan, labels):
        _write(book, spec, rows)
        if spec.title not in {labels.summary, labels.categories}:
            keys += [str(row[0]) for row in rows]
    meta = cast("WriteOnly", book.create_sheet(META_SHEET))
    meta.sheet_state = "veryHidden"
    meta.protection.sheet = True
    listed = [labels.stay, *plan.categories()]
    pairs = labels.meta_rows(plan.plan_id, fingerprint(plan.plan_id, keys))
    for index in range(max(len(pairs), len(listed))):
        key, value = pairs[index] if index < len(pairs) else (None, None)
        choice = listed[index] if index < len(listed) else None
        meta.append([key, value, None, choice])
    book.save(target)


def _sheets(plan: ClassifyPlan, labels: Labels) -> list[tuple[_Spec, list[Row]]]:
    """The four visible sheets and their rows, in the workbook's order.

    Args:
        plan: The plan.
        labels: The translated names and values.

    Returns:
        Each sheet with its rows.
    """
    last = len(plan.categories()) + 1
    choices = f"'{META_SHEET}'!${META_LIST_COLUMN}$1:${META_LIST_COLUMN}${last}"
    category = CategoryColumn
    event = EventColumn
    return [
        (
            _Spec(
                labels.summary,
                summary_headers(labels),
                frozenset({_SUMMARY_NOTES}),
                widths=_SUMMARY_WIDTHS,
                filtered=False,
            ),
            summary_rows(plan),
        ),
        (
            _Spec(
                labels.categories,
                content.category_headers(),
                frozenset({category.RENAME, category.CONFIRM, category.NOTES}),
                {
                    category.RENAME: choices,
                    category.CONFIRM: f'"{labels.yes},{labels.no}"',
                },
            ),
            content.category_rows(plan),
        ),
        (
            _Spec(
                labels.events,
                content.event_headers(),
                frozenset({event.NAME, event.CATEGORY, event.NOTES}),
                {event.CATEGORY: choices},
                {event.SHARE: _PERCENT},
            ),
            content.event_rows(plan, labels),
        ),
        (
            _Spec(
                labels.files,
                content.file_headers(),
                frozenset({FileColumn.FINAL, FileColumn.NOTES}),
                {FileColumn.FINAL: f'"{labels.stay}"'},
            ),
            content.file_rows(plan, labels),
        ),
    ]


def _write(book: Workbook, spec: _Spec, rows: list[Row]) -> None:
    """Add one protected sheet, its header row, its rows and its drop-downs.

    Args:
        book: The workbook, write-only.
        spec: The sheet.
        rows: Its rows.
    """
    created: WriteOnlyWorksheet = book.create_sheet(spec.title)
    sheet = cast("WriteOnly", created)
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False  # False: filtering is allowed
    sheet.freeze_panes = "A2"
    width = max(len(spec.headers), len(spec.widths))
    last = len(rows) + 1
    for column in range(1, width + 1):
        size = spec.widths[column - 1] if column <= len(spec.widths) else _WIDTH
        sheet.column_dimensions[get_column_letter(column)].width = size
    for column, formula in spec.lists.items():
        validation = DataValidation(
            type="list", formula1=formula, allow_blank=True, showErrorMessage=False
        )
        letter = get_column_letter(column)
        validation.add(f"{letter}2:{letter}{last}")
        sheet.data_validations.append(validation)
    if spec.filtered:
        sheet.auto_filter.ref = f"A1:{get_column_letter(width)}{last}"
    sheet.append([header_cell(created, value) for value in spec.headers])
    for row in rows:
        sheet.append(
            [
                editable_cell(created, value)
                if column in spec.editable
                else locked_cell(created, value, spec.formats.get(column))
                for column, value in enumerate(padded(row, max(spec.editable)), start=1)
            ]
        )
