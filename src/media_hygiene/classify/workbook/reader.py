"""Read back the edits of `classify.xlsx`, and refuse a workbook that was damaged.

The hidden `_meta` sheet says which plan the workbook belongs to and holds a fingerprint
of its locked keys; `structure` compares every locked cell with the plan: a sheet
renamed, a row deleted or a workbook of another run are refused. Every problem names
its cell.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Final
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from media_hygiene.classify.workbook.edits import Edits
from media_hygiene.classify.workbook.sheet_edits import (
    Problems,
    at,
    category_edits,
    event_edits,
    file_edits,
)
from media_hygiene.classify.workbook.sheets import (
    META_SHEET,
    WORKBOOK_FORMAT,
    Labels,
    MetaKey,
)
from media_hygiene.classify.workbook.structure import (
    ReadBook,
    check_structure,
    undo_tip,
)
from media_hygiene.classify.workbook.validation import text_of
from media_hygiene.errors import WorkbookError
from media_hygiene.i18n import _

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook

    from media_hygiene.classify.plan_file import ClassifyPlan
    from media_hygiene.classify.workbook.sheet_edits import Table


_SHOWN_PROBLEMS: Final = 10


def read_edits(path: Path, plan: ClassifyPlan) -> Edits:
    """Read the editable cells of a workbook written for `plan`.

    Args:
        path: The `.xlsx` file.
        plan: Its plan.

    Returns:
        What the user changed.

    Raises:
        WorkbookError: The workbook is not the plan's, was damaged, or holds a value
            that cannot be used.
    """
    read = _read(path)
    if read.meta[MetaKey.PLAN_ID] != plan.plan_id:
        raise WorkbookError(_("This workbook belongs to another classify run."))
    check_structure(read, plan)
    labels = read.labels
    categories, events, files = (
        read.tables[name][1:]
        for name in (labels.categories, labels.events, labels.files)
    )
    problems = Problems(labels)
    edits = Edits(
        files=file_edits(files, problems),
        events=event_edits(events, problems),
        categories=category_edits(categories, problems),
    )
    if problems.found:
        raise WorkbookError(
            _("The workbook holds values that cannot be used:")
            + "\n"
            + "\n".join(problems.found[:_SHOWN_PROBLEMS]),
            _("Correct these cells, save the workbook, then run 'sort' again."),
        )
    return edits


def read_plan_id(path: Path) -> str:
    """Tell which classify run a workbook belongs to.

    Args:
        path: The `.xlsx` file.

    Returns:
        The plan id its `_meta` sheet records.

    Raises:
        WorkbookError: It is not a classify workbook.
    """
    book = _open(path)
    try:
        return _meta(book)[MetaKey.PLAN_ID]
    finally:
        book.close()


def _open(path: Path) -> Workbook:
    """Open a workbook to read its values.

    Args:
        path: The `.xlsx` file.

    Returns:
        The workbook, read-only.

    Raises:
        WorkbookError: It cannot be opened.
    """
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException, KeyError) as exc:
        raise WorkbookError(
            _("The workbook {path} cannot be opened: {reason}.").format(
                path=path.name, reason=exc
            )
        ) from exc


def _read(path: Path) -> ReadBook:
    """Read every sheet of a workbook, as values.

    Args:
        path: The `.xlsx` file.

    Returns:
        Its sheet names, rows and `_meta`.

    Raises:
        WorkbookError: It is not a classify workbook, or a sheet is missing.
    """
    book = _open(path)
    try:
        meta = _meta(book)
        labels = Labels.from_meta(meta)
        names = (labels.summary, labels.categories, labels.events, labels.files)
        missing = [name for name in names if name not in book.sheetnames]
        if missing:
            raise WorkbookError(
                _("The sheet '{name}' was renamed or deleted.").format(name=missing[0]),
                undo_tip(),
            )
        tables = {name: _table(book, name, first=1) for name in names}
        return ReadBook(tuple(book.sheetnames), tables, meta, labels)
    finally:
        book.close()


def _meta(book: Workbook) -> dict[MetaKey, str]:
    """Read the hidden `_meta` sheet.

    Args:
        book: The workbook.

    Returns:
        Its values.

    Raises:
        WorkbookError: It is missing, or of another format.
    """
    if META_SHEET not in book.sheetnames:
        raise WorkbookError(_("This is not a classify workbook, or it was damaged."))
    rows = _table(book, META_SHEET, first=1)
    values = {key: text_of(at(rows[key - 1], 2)) for key in MetaKey if key <= len(rows)}
    if values.get(MetaKey.FORMAT) != str(WORKBOOK_FORMAT) or len(values) < len(MetaKey):
        raise WorkbookError(_("This is not a classify workbook, or it was damaged."))
    return values


def _table(book: Workbook, name: str, first: int = 2) -> Table:
    """The rows of a sheet, as values.

    Args:
        book: The workbook, read-only.
        name: The sheet.
        first: The first row read (2: below the headers).

    Returns:
        The rows.
    """
    sheet = book[name]
    return [tuple(row) for row in sheet.iter_rows(min_row=first, values_only=True)]
