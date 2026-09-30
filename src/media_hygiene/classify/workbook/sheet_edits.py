"""The editable cells of each sheet, checked one by one: problems name their cell."""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.utils import get_column_letter

from media_hygiene.classify.workbook.edits import CategoryEdit, EventEdit
from media_hygiene.classify.workbook.sheets import (
    CategoryColumn,
    EventColumn,
    FileColumn,
)
from media_hygiene.classify.workbook.validation import (
    choice_of,
    folder_problem,
    text_of,
)
from media_hygiene.i18n import _

if TYPE_CHECKING:
    from collections.abc import Sequence

    from media_hygiene.classify.workbook.edits import Choice
    from media_hygiene.classify.workbook.sheets import Labels

type Table = list[Sequence[object]]


class Problems:
    """The problems found so far, each with its cell."""

    def __init__(self, labels: Labels) -> None:
        """Start with none.

        Args:
            labels: The workbook's values.
        """
        self.labels = labels
        self.found: list[str] = []

    def choice(self, text: str, cell: tuple[str, int, int]) -> Choice:
        """Check a folder or "(stay where it is)".

        Args:
            text: The cell's text, not empty.
            cell: Sheet, row and column, for the message.

        Returns:
            The value; an empty string when it cannot be used.
        """
        if text == self.labels.stay:
            return choice_of(text, self.labels)
        problem = folder_problem(text)
        if problem is None:
            return choice_of(text, self.labels)
        sheet, row, column = cell
        self.found.append(f"{sheet}!{get_column_letter(column)}{row}: {problem}")
        return ""


def at(row: Sequence[object], column: int) -> object:
    """A cell of a row, None past its end.

    Args:
        row: The row.
        column: 1-based.

    Returns:
        The value.
    """
    return row[column - 1] if column <= len(row) else None


def file_edits(
    table: Table, problems: Problems
) -> tuple[dict[str, Choice], dict[str, str]]:
    """The final folders typed on the Files sheet, and their cells.

    Args:
        table: Its rows.
        problems: Where problems go.

    Returns:
        Row id → folder or "stay"; row id → its cell (`Files!K7`), for the messages.
    """
    edits: dict[str, Choice] = {}
    cells: dict[str, str] = {}
    sheet, column = problems.labels.files, get_column_letter(FileColumn.FINAL)
    for number, row in enumerate(table, start=2):
        text = text_of(at(row, FileColumn.FINAL))
        if text:
            row_id = text_of(row[0])
            edits[row_id] = problems.choice(text, (sheet, number, FileColumn.FINAL))
            cells[row_id] = f"{sheet}!{column}{number}"
    kept = {key: value for key, value in edits.items() if value}
    return kept, {key: cells[key] for key in kept}


def event_edits(table: Table, problems: Problems) -> dict[str, EventEdit]:
    """The names and categories typed on the Events sheet.

    Args:
        table: Its rows.
        problems: Where problems go.

    Returns:
        Event id → its edit.
    """
    edits: dict[str, EventEdit] = {}
    sheet = problems.labels.events
    for number, row in enumerate(table, start=2):
        name = text_of(at(row, EventColumn.NAME))
        category = text_of(at(row, EventColumn.CATEGORY))
        if name:
            checked = problems.choice(name, (sheet, number, EventColumn.NAME))
            name = checked if isinstance(checked, str) else ""
        chosen = (
            problems.choice(category, (sheet, number, EventColumn.CATEGORY))
            if category
            else ""
        )
        if name or chosen:
            edits[text_of(row[0])] = EventEdit(name, chosen)
    return edits


def category_edits(table: Table, problems: Problems) -> dict[str, CategoryEdit]:
    """The new names and confirmations of the Categories sheet.

    Args:
        table: Its rows.
        problems: Where problems go.

    Returns:
        Category → its edit.
    """
    edits: dict[str, CategoryEdit] = {}
    labels = problems.labels
    for number, row in enumerate(table, start=2):
        rename = text_of(at(row, CategoryColumn.RENAME))
        confirm = text_of(at(row, CategoryColumn.CONFIRM))
        cell = (labels.categories, number, CategoryColumn.RENAME)
        chosen = problems.choice(rename, cell) if rename else ""
        if confirm not in {"", labels.yes, labels.no}:
            column = get_column_letter(CategoryColumn.CONFIRM)
            problems.found.append(
                f"{labels.categories}!{column}{number}: "
                + _("'{text}': choose {yes} or {no}.").format(
                    text=confirm, yes=labels.yes, no=labels.no
                )
            )
        if chosen or confirm == labels.yes:
            edits[text_of(row[0])] = CategoryEdit(chosen, confirm == labels.yes)
    return edits
